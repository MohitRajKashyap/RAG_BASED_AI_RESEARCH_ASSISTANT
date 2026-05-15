import logging
from pathlib import Path
from typing import List

from langchain_core.documents import Document

logger = logging.getLogger(__name__)


def load_pdf(file_path: str) -> List[Document]:
    docs: List[Document] = []
    try:
        import pdfplumber

        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(
                        Document(
                            page_content=text,
                            metadata={"source": file_path, "page": i + 1},
                        )
                    )
    except Exception:
        from PyPDF2 import PdfReader

        reader = PdfReader(file_path)
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                docs.append(
                    Document(
                        page_content=text,
                        metadata={"source": file_path, "page": i + 1},
                    )
                )
    return docs


def load_docx(file_path: str) -> List[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    return [Document(page_content=text, metadata={"source": file_path, "page": 1})] if text else []


def load_txt(file_path: str) -> List[Document]:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    return [Document(page_content=text, metadata={"source": file_path, "page": 1})] if text.strip() else []


def load_document(file_path: str, file_type: str) -> List[Document]:
    path = Path(file_path)
    ext = path.suffix.lower()
    logger.info("Loading document %s (%s)", file_path, file_type)
    if ext == ".pdf":
        return load_pdf(file_path)
    if ext in (".docx", ".doc"):
        return load_docx(file_path)
    if ext == ".txt":
        return load_txt(file_path)
    raise ValueError(f"Unsupported file type: {ext}")
